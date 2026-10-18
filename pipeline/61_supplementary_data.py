# Additional files 2-5 of the article as Excel workbooks
# additional_file_2_tables.xlsx: Supplementary Tables S1-S7 (source_data/supplementary_table*.tsv)
# additional_file_3_age_effects.xlsx: the per-gene age effects of every unit in the three layers (script 10)
# additional_file_4_clock_weights.xlsx: the clock weights of the mouse and human proteome models, with and without the
# m6A features (script 52, models of all cohorts), and the gene-level model predictions (script 50)
# additional_file_5_clock_scores.xlsx: the clock scores of every sample (scripts 52 and 53) and the individual m6A
# index (script 54)
# the README of each workbook comes from docs/supplementary_data_columns.tsv (one row per sheet, then one per column;
# sheet "*" holds the columns shared by all sheets, "* tables" those of the tables); a column without a description, or
# a description without its column, stops the script
# reads output/10, 50, 52, 53 and 54, source_data/ and the column table; writes output/61_supplementary_data/
import time

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import utils

OUT = utils.OUT / "61_supplementary_data"
CHUNK = 100_000  # rows converted and written at a time
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
FILES = {
    "2": "additional_file_2_tables",
    "3": "additional_file_3_age_effects",
    "4": "additional_file_4_clock_weights",
    "5": "additional_file_5_clock_scores",
}

described = utils.read(utils.ROOT / "docs" / "supplementary_data_columns.tsv", keep_default_na=False, dtype=str)
entries = described[described.column == ""]  # one row per sheet, in workbook order
columns_of = {sheet: rows.set_index("column")[["meaning", "unit"]] for sheet, rows in described.groupby("sheet")}


def weights(species):
    runs = utils.read(utils.OUT / "52_proteome_clock" / "runs" / f"{species}_protein__all.parquet")
    runs = runs.rename(columns={"model": "weight", "without_m6a": "weight_without_m6a"})
    return runs[["symbol", "weight", "weight_without_m6a"]].rename_axis("gene_id").reset_index()


def read_table(file):
    if file.startswith("52_proteome_clock/runs/"):
        return weights(file.split("/")[-1].split("_")[0])
    path = utils.ROOT / file if file.startswith("source_data/") else utils.OUT / file
    return utils.read(path, keep_default_na=False, na_values=[""], low_memory=False)  # only empty fields are missing


def readme_rows(book, columns):
    # a row per sheet and, below it, a row per column; every column of every table must be described, and every
    # description must belong to a column
    rows = []
    for entry in entries[entries.workbook == book].itertuples():
        script, source = f"pipeline/{entry.file.split('/')[0]}.py", f"output/{entry.file}"
        if entry.file.startswith("source_data/"):
            script, source = "pipeline/60_source_data.py", entry.file
        rows.append((entry.sheet, None, entry.meaning, None, script, source))
        own = columns_of[entry.sheet].drop(index="")
        shared = [columns_of["*"]] + ([columns_of["* tables"]] if entry.file.startswith("source_data/") else [])
        known = pd.concat([own, *shared])
        known = known[~known.index.duplicated()]  # the sheet's own meaning first, then the shared ones
        missing = [column for column in columns[entry.sheet] if column not in known.index]
        extra = [column for column in own.index if column not in columns[entry.sheet]]
        assert not missing and not extra, f"{entry.sheet}: not described {missing}, described but absent {extra}"
        rows += [(entry.sheet, column, *known.loc[column], None, None) for column in columns[entry.sheet]]
    return rows


def add_sheet(wb, name, table):
    # header row frozen, bold and filtered; numbers stay numbers, missing values stay empty; columns sized from the head
    ws = wb.create_sheet(name)
    ws.freeze_panes = "A2"
    header = []
    for i, column in enumerate(table, 1):
        lengths = table[column].head(1000).dropna().astype(str).str.len()
        text = min(max(len(column), lengths.max() if len(lengths) else 0), 60)
        numeric = pd.api.types.is_numeric_dtype(table[column])
        ws.column_dimensions[get_column_letter(i)].width = max(len(column), 10) + 2 if numeric else text + 2
        header.append(WriteOnlyCell(ws, value=column))
        header[-1].font = BOLD
    ws.append(header)
    for start in range(0, len(table), CHUNK):
        part = table.iloc[start : start + CHUNK]
        for row in part.astype(object).where(part.notna(), None).itertuples(index=False, name=None):
            ws.append(row)
    ws.auto_filter.ref = f"A1:{get_column_letter(table.shape[1])}{len(table) + 1}"


def add_readme(wb, rows):
    ws = wb.create_sheet("README")
    ws.freeze_panes = "A2"
    for letter, width in zip("ABCDEF", (32, 30, 110, 44, 40, 70), strict=True):
        ws.column_dimensions[letter].width = width

    def cells(values, bold):
        out = [WriteOnlyCell(ws, value=value) for value in values]
        for cell in out:
            cell.alignment = WRAP
            cell.font = BOLD if bold else Font()
        return out

    ws.append(cells(["Sheet", "Column", "Meaning", "Measurement unit", "Source script", "Output file"], True))
    for row in rows:
        ws.append(cells(row, row[1] is None))
    ws.auto_filter.ref = f"A1:F{len(rows) + 1}"


assert entries.sheet.is_unique and entries.sheet.str.len().max() <= 31, "sheet names: unique, 31 characters at most"
start = time.time()
OUT.mkdir(parents=True, exist_ok=True)
for book, name in FILES.items():
    files = dict(entries[entries.workbook == book][["sheet", "file"]].itertuples(index=False))
    tables = {sheet: read_table(file) for sheet, file in files.items()}
    wb = Workbook(write_only=True)
    add_readme(wb, readme_rows(book, {sheet: list(table) for sheet, table in tables.items()}))
    for sheet, table in tables.items():
        add_sheet(wb, sheet, table)
    path = OUT / f"{name}.xlsx"
    wb.save(path)
    print(f"{path.name}: {path.stat().st_size / 1e6:.0f} MB, {time.time() - start:.0f} s", flush=True)
